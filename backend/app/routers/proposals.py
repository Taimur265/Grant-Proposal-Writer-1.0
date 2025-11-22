"""Proposal API routes."""

from typing import Optional, List
from uuid import UUID
from datetime import datetime
import io

from fastapi import APIRouter, Depends, HTTPException, status, Query, BackgroundTasks
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func

from app.database import get_db
from app.models.project import Project
from app.models.proposal import Proposal, ProposalSection, ProposalStatus, ProposalFeedback
from app.models.user import User
from app.schemas.proposal import (
    ProposalCreate, ProposalUpdate, ProposalResponse, ProposalListResponse,
    ProposalGenerate, ProposalSectionUpdate, ProposalSectionResponse,
    ProposalExport, ProposalFeedbackCreate, ProposalFeedbackResponse, ComplianceCheck
)
from app.routers.auth import get_current_user
from app.services.proposal_generator import ProposalGenerator
from app.services.ai_service import AIService
from app.config import settings

router = APIRouter(prefix="/proposals", tags=["Proposals"])


@router.post("/generate", response_model=ProposalResponse, status_code=status.HTTP_201_CREATED)
async def generate_proposal(
    generation_request: ProposalGenerate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Generate a new proposal using AI."""
    # Verify project access
    project = await db.get(Project, generation_request.project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    if project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized")

    # Initialize AI service with requested provider/model
    ai_service = AIService(
        provider=generation_request.ai_provider,
        model=generation_request.model,
    )

    # Generate proposal
    generator = ProposalGenerator(db, ai_service)
    proposal = await generator.generate_proposal(
        project_id=generation_request.project_id,
        user_id=current_user.id,
        title=generation_request.title,
        tone=generation_request.tone,
        focus_areas=generation_request.focus_areas,
        custom_instructions=generation_request.custom_instructions,
        include_sections=generation_request.include_sections,
        max_words_per_section=generation_request.max_words_per_section,
    )

    # Load sections for response
    result = await db.execute(
        select(ProposalSection)
        .where(ProposalSection.proposal_id == proposal.id)
        .order_by(ProposalSection.section_order)
    )
    proposal.sections = result.scalars().all()

    return proposal


@router.get("/", response_model=ProposalListResponse)
async def list_proposals(
    project_id: Optional[UUID] = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    status_filter: Optional[ProposalStatus] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List proposals for a project or all user's proposals."""
    # Build base query - join with projects to filter by owner
    query = (
        select(Proposal)
        .join(Project)
        .where(Project.owner_id == current_user.id)
    )

    if project_id:
        query = query.where(Proposal.project_id == project_id)
    if status_filter:
        query = query.where(Proposal.status == status_filter)

    # Get total count
    count_query = select(func.count()).select_from(query.subquery())
    total = (await db.execute(count_query)).scalar()

    # Apply pagination
    query = query.order_by(Proposal.updated_at.desc())
    query = query.offset((page - 1) * page_size).limit(page_size)

    result = await db.execute(query)
    proposals = result.scalars().all()

    # Load sections for each proposal
    for proposal in proposals:
        sections_result = await db.execute(
            select(ProposalSection)
            .where(ProposalSection.proposal_id == proposal.id)
            .order_by(ProposalSection.section_order)
        )
        proposal.sections = sections_result.scalars().all()

    return ProposalListResponse(
        items=proposals,
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get("/{proposal_id}", response_model=ProposalResponse)
async def get_proposal(
    proposal_id: UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Get a specific proposal with all sections."""
    proposal = await db.get(Proposal, proposal_id)
    if not proposal:
        raise HTTPException(status_code=404, detail="Proposal not found")

    # Verify access
    project = await db.get(Project, proposal.project_id)
    if project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized")

    # Load sections
    result = await db.execute(
        select(ProposalSection)
        .where(ProposalSection.proposal_id == proposal_id)
        .order_by(ProposalSection.section_order)
    )
    proposal.sections = result.scalars().all()

    return proposal


@router.patch("/{proposal_id}", response_model=ProposalResponse)
async def update_proposal(
    proposal_id: UUID,
    proposal_update: ProposalUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Update proposal metadata."""
    proposal = await db.get(Proposal, proposal_id)
    if not proposal:
        raise HTTPException(status_code=404, detail="Proposal not found")

    # Verify access
    project = await db.get(Project, proposal.project_id)
    if project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized")

    for field, value in proposal_update.model_dump(exclude_unset=True).items():
        setattr(proposal, field, value)

    await db.flush()
    await db.refresh(proposal)
    return proposal


@router.delete("/{proposal_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_proposal(
    proposal_id: UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Delete a proposal."""
    proposal = await db.get(Proposal, proposal_id)
    if not proposal:
        raise HTTPException(status_code=404, detail="Proposal not found")

    # Verify access
    project = await db.get(Project, proposal.project_id)
    if project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized")

    await db.delete(proposal)


@router.patch("/{proposal_id}/sections/{section_id}", response_model=ProposalSectionResponse)
async def update_section(
    proposal_id: UUID,
    section_id: UUID,
    section_update: ProposalSectionUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Update a proposal section."""
    section = await db.get(ProposalSection, section_id)
    if not section or section.proposal_id != proposal_id:
        raise HTTPException(status_code=404, detail="Section not found")

    # Verify access
    proposal = await db.get(Proposal, proposal_id)
    project = await db.get(Project, proposal.project_id)
    if project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized")

    if section_update.content is not None:
        section.content = section_update.content
        section.word_count = len(section_update.content.split())
    if section_update.is_complete is not None:
        section.is_complete = section_update.is_complete

    await db.flush()
    await db.refresh(section)
    return section


@router.post("/{proposal_id}/sections/{section_id}/regenerate", response_model=ProposalSectionResponse)
async def regenerate_section(
    proposal_id: UUID,
    section_id: UUID,
    feedback: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Regenerate a proposal section with optional feedback."""
    section = await db.get(ProposalSection, section_id)
    if not section or section.proposal_id != proposal_id:
        raise HTTPException(status_code=404, detail="Section not found")

    # Verify access
    proposal = await db.get(Proposal, proposal_id)
    project = await db.get(Project, proposal.project_id)
    if project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized")

    # Regenerate using proposal generator
    generator = ProposalGenerator(db)
    updated_section = await generator.regenerate_section(
        proposal_id=proposal_id,
        section_id=section_id,
        feedback=feedback,
    )

    return updated_section


@router.post("/{proposal_id}/check-compliance", response_model=ComplianceCheck)
async def check_compliance(
    proposal_id: UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Check proposal compliance against guidelines."""
    proposal = await db.get(Proposal, proposal_id)
    if not proposal:
        raise HTTPException(status_code=404, detail="Proposal not found")

    # Verify access
    project = await db.get(Project, proposal.project_id)
    if project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized")

    # Get guidelines documents
    from app.models.document import Document, DocumentCategory
    result = await db.execute(
        select(Document).where(
            Document.project_id == proposal.project_id,
            Document.category == DocumentCategory.GUIDELINES,
            Document.processing_status == "completed",
        )
    )
    guidelines_docs = result.scalars().all()

    if not guidelines_docs:
        raise HTTPException(
            status_code=400,
            detail="No processed guidelines documents found for compliance check",
        )

    # Combine guidelines content
    guidelines_content = "\n\n".join([
        doc.extracted_text for doc in guidelines_docs if doc.extracted_text
    ])

    # Run compliance check
    ai_service = AIService()
    compliance_result = await ai_service.check_compliance(
        proposal.full_content or "",
        guidelines_content,
    )

    # Update proposal compliance score if available
    if isinstance(compliance_result.get("score"), (int, float)):
        proposal.compliance_score = compliance_result["score"]
        await db.flush()

    return ComplianceCheck(
        overall_score=compliance_result.get("score", 0),
        section_scores=compliance_result.get("section_scores", {}),
        missing_requirements=compliance_result.get("missing_requirements", []),
        suggestions=compliance_result.get("suggestions", []),
        word_count_status=compliance_result.get("word_count_status", {}),
    )


@router.post("/{proposal_id}/export")
async def export_proposal(
    proposal_id: UUID,
    export_request: ProposalExport,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Export proposal to specified format."""
    proposal = await db.get(Proposal, proposal_id)
    if not proposal:
        raise HTTPException(status_code=404, detail="Proposal not found")

    # Verify access
    project = await db.get(Project, proposal.project_id)
    if project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized")

    # Load sections
    result = await db.execute(
        select(ProposalSection)
        .where(ProposalSection.proposal_id == proposal_id)
        .order_by(ProposalSection.section_order)
    )
    sections = result.scalars().all()

    # Generate content based on format
    format_type = export_request.format

    if format_type == "txt":
        content = _generate_txt(proposal, sections, export_request.include_metadata)
        media_type = "text/plain"
        filename = f"{proposal.title}.txt"
    elif format_type == "md":
        content = _generate_markdown(proposal, sections, export_request.include_metadata)
        media_type = "text/markdown"
        filename = f"{proposal.title}.md"
    elif format_type == "html":
        content = _generate_html(proposal, sections, export_request.include_metadata)
        media_type = "text/html"
        filename = f"{proposal.title}.html"
    elif format_type == "docx":
        content = await _generate_docx(proposal, sections, export_request.include_metadata)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        filename = f"{proposal.title}.docx"
    elif format_type == "pdf":
        raise HTTPException(status_code=501, detail="PDF export not yet implemented")
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported format: {format_type}")

    # Update export timestamp
    proposal.last_exported_at = datetime.utcnow()
    proposal.export_format = format_type
    await db.flush()

    # Return file
    if isinstance(content, bytes):
        return StreamingResponse(
            io.BytesIO(content),
            media_type=media_type,
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    else:
        return StreamingResponse(
            io.BytesIO(content.encode('utf-8')),
            media_type=media_type,
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )


def _generate_txt(proposal: Proposal, sections: List[ProposalSection], include_metadata: bool) -> str:
    """Generate plain text export."""
    lines = []

    if include_metadata:
        lines.append(f"Title: {proposal.title}")
        lines.append(f"Status: {proposal.status.value}")
        lines.append(f"Created: {proposal.created_at.strftime('%Y-%m-%d')}")
        if proposal.compliance_score:
            lines.append(f"Compliance Score: {proposal.compliance_score:.1f}%")
        lines.append("")
        lines.append("=" * 60)
        lines.append("")

    if proposal.executive_summary:
        lines.append("EXECUTIVE SUMMARY")
        lines.append("-" * 40)
        lines.append(proposal.executive_summary)
        lines.append("")

    for section in sections:
        lines.append(section.section_name.upper())
        lines.append("-" * 40)
        lines.append(section.content or "")
        lines.append("")

    return "\n".join(lines)


def _generate_markdown(proposal: Proposal, sections: List[ProposalSection], include_metadata: bool) -> str:
    """Generate Markdown export."""
    lines = []

    lines.append(f"# {proposal.title}")
    lines.append("")

    if include_metadata:
        lines.append("---")
        lines.append(f"**Status:** {proposal.status.value}")
        lines.append(f"**Created:** {proposal.created_at.strftime('%Y-%m-%d')}")
        if proposal.compliance_score:
            lines.append(f"**Compliance Score:** {proposal.compliance_score:.1f}%")
        lines.append("---")
        lines.append("")

    if proposal.executive_summary:
        lines.append("## Executive Summary")
        lines.append("")
        lines.append(proposal.executive_summary)
        lines.append("")

    for section in sections:
        lines.append(f"## {section.section_name}")
        lines.append("")
        lines.append(section.content or "")
        lines.append("")

    return "\n".join(lines)


def _generate_html(proposal: Proposal, sections: List[ProposalSection], include_metadata: bool) -> str:
    """Generate HTML export."""
    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{proposal.title}</title>
    <style>
        body {{ font-family: 'Georgia', serif; max-width: 800px; margin: 0 auto; padding: 40px; line-height: 1.6; }}
        h1 {{ color: #2c3e50; border-bottom: 2px solid #3498db; padding-bottom: 10px; }}
        h2 {{ color: #34495e; margin-top: 30px; }}
        .metadata {{ background: #f8f9fa; padding: 15px; border-radius: 5px; margin-bottom: 30px; }}
        .metadata p {{ margin: 5px 0; }}
        .section {{ margin-bottom: 30px; }}
    </style>
</head>
<body>
    <h1>{proposal.title}</h1>
"""

    if include_metadata:
        html += f"""    <div class="metadata">
        <p><strong>Status:</strong> {proposal.status.value}</p>
        <p><strong>Created:</strong> {proposal.created_at.strftime('%Y-%m-%d')}</p>
"""
        if proposal.compliance_score:
            html += f"        <p><strong>Compliance Score:</strong> {proposal.compliance_score:.1f}%</p>\n"
        html += "    </div>\n"

    if proposal.executive_summary:
        html += f"""    <div class="section">
        <h2>Executive Summary</h2>
        <p>{proposal.executive_summary.replace(chr(10), '</p><p>')}</p>
    </div>
"""

    for section in sections:
        content = (section.content or "").replace("\n", "</p><p>")
        html += f"""    <div class="section">
        <h2>{section.section_name}</h2>
        <p>{content}</p>
    </div>
"""

    html += """</body>
</html>"""

    return html


async def _generate_docx(proposal: Proposal, sections: List[ProposalSection], include_metadata: bool) -> bytes:
    """Generate DOCX export."""
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    title = doc.add_heading(proposal.title, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if include_metadata:
        meta_para = doc.add_paragraph()
        meta_para.add_run(f"Status: {proposal.status.value}\n").bold = True
        meta_para.add_run(f"Created: {proposal.created_at.strftime('%Y-%m-%d')}\n")
        if proposal.compliance_score:
            meta_para.add_run(f"Compliance Score: {proposal.compliance_score:.1f}%")
        doc.add_paragraph()

    if proposal.executive_summary:
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(proposal.executive_summary)

    for section in sections:
        doc.add_heading(section.section_name, level=1)
        if section.content:
            doc.add_paragraph(section.content)

    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


@router.post("/{proposal_id}/feedback", response_model=ProposalFeedbackResponse, status_code=status.HTTP_201_CREATED)
async def add_feedback(
    proposal_id: UUID,
    feedback_data: ProposalFeedbackCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Add feedback to a proposal."""
    proposal = await db.get(Proposal, proposal_id)
    if not proposal:
        raise HTTPException(status_code=404, detail="Proposal not found")

    # Verify access
    project = await db.get(Project, proposal.project_id)
    if project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized")

    feedback = ProposalFeedback(
        proposal_id=proposal_id,
        section_id=feedback_data.section_id,
        feedback_type=feedback_data.feedback_type,
        content=feedback_data.content,
        author_id=current_user.id,
    )
    db.add(feedback)
    await db.flush()
    await db.refresh(feedback)

    return feedback


@router.get("/{proposal_id}/feedback", response_model=List[ProposalFeedbackResponse])
async def list_feedback(
    proposal_id: UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List all feedback for a proposal."""
    proposal = await db.get(Proposal, proposal_id)
    if not proposal:
        raise HTTPException(status_code=404, detail="Proposal not found")

    # Verify access
    project = await db.get(Project, proposal.project_id)
    if project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized")

    result = await db.execute(
        select(ProposalFeedback)
        .where(ProposalFeedback.proposal_id == proposal_id)
        .order_by(ProposalFeedback.created_at.desc())
    )

    return result.scalars().all()
